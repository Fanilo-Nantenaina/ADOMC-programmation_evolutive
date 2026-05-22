import io
from datetime import datetime
from typing import Any, Dict, List

import pandas as pd
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import DataBarRule, Rule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.schemas.inputs import ReportExportRequest

router = APIRouter()


COLOR_PRIMARY = "1E3A5F"
COLOR_PRIMARY_LIGHT = "E8EEF5"
COLOR_TEXT = "1A1A1A"
COLOR_MUTED = "6B7280"
COLOR_ZEBRA = "F8FAFC"
COLOR_BORDER = "D1D5DB"

FONT_TITLE = Font(name="Calibri", size=18, bold=True, color=COLOR_PRIMARY)
FONT_SECTION = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color=COLOR_PRIMARY)
FONT_BODY = Font(name="Calibri", size=10, color=COLOR_TEXT)
FONT_MONO = Font(name="Consolas", size=10, color=COLOR_TEXT)
FONT_CAPTION = Font(name="Calibri", size=9, italic=True, color=COLOR_MUTED)

FILL_PRIMARY = PatternFill("solid", fgColor=COLOR_PRIMARY)
FILL_HEADER = PatternFill("solid", fgColor=COLOR_PRIMARY_LIGHT)
FILL_ZEBRA = PatternFill("solid", fgColor=COLOR_ZEBRA)

THIN = Side(style="thin", color=COLOR_BORDER)
BORDER_CELL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center", indent=1)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center", indent=1)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")


def _style_section_title(ws, row: int, text: str, span: int) -> None:
    """Insert a section title row with primary fill."""
    ws.cell(row=row, column=1, value=text).font = FONT_SECTION
    ws.cell(row=row, column=1).fill = FILL_PRIMARY
    ws.cell(row=row, column=1).alignment = Alignment(
        horizontal="left", vertical="center", indent=1
    )
    ws.merge_cells(start_row=row, end_row=row, start_column=1, end_column=span)
    ws.row_dimensions[row].height = 26


def _style_header_row(ws, row: int, n_cols: int) -> None:
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.border = BORDER_CELL
        cell.alignment = ALIGN_CENTER
    ws.row_dimensions[row].height = 22


def _apply_zebra(ws, start_row: int, end_row: int, n_cols: int) -> None:
    for r in range(start_row, end_row + 1):
        if (r - start_row) % 2 == 1:
            for c in range(1, n_cols + 1):
                ws.cell(row=r, column=c).fill = FILL_ZEBRA


def _set_col_widths(ws, widths: List[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _build_config_sheet(wb: Workbook, req: ReportExportRequest) -> None:
    ws = wb.create_sheet("Configuration")

    ws["A1"] = "Rapport SID — Système d'Aide à la Décision"
    ws["A1"].font = FONT_TITLE
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A1:B1")
    ws["A2"] = "Optimisation évolutive multi-objectif de portefeuille"
    ws["A2"].font = FONT_CAPTION
    ws.merge_cells("A2:B2")

    _style_section_title(ws, 4, "Paramètres de la simulation", span=2)

    rows: List[List[Any]] = [
        ("Date de simulation", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Algorithme", req.config.algorithm.upper()),
        ("Population", req.config.pop_size),
        ("Générations", req.config.n_gen),
        ("Seed", req.config.seed),
        ("Priorité Rendement", f"{req.config.w_return:.0f} %"),
        ("Priorité Risque", f"{100 - req.config.w_return:.0f} %"),
        ("Taux sans risque", f"{req.config.risk_free_rate * 100:.2f} %"),
        (
            "Contrainte σ_max",
            (
                f"{req.config.max_risk_pct * 100:.0f} %"
                if req.config.max_risk_pct is not None
                else "Aucune"
            ),
        ),
        ("Nombre d'actifs", len(req.config.assets)),
    ]

    _style_header_row(ws, 5, 2)
    ws.cell(row=5, column=1, value="Paramètre Décisionnel")
    ws.cell(row=5, column=2, value="Valeur")
    _style_header_row(ws, 5, 2)

    for i, (k, v) in enumerate(rows, start=6):
        ws.cell(row=i, column=1, value=k).font = FONT_BODY
        ws.cell(row=i, column=1).alignment = ALIGN_LEFT
        ws.cell(row=i, column=1).border = BORDER_CELL

        ws.cell(row=i, column=2, value=v).font = FONT_MONO
        ws.cell(row=i, column=2).alignment = ALIGN_RIGHT
        ws.cell(row=i, column=2).border = BORDER_CELL

    _apply_zebra(ws, 6, 5 + len(rows), 2)
    _set_col_widths(ws, [30, 32])

    _style_section_title(ws, 5 + len(rows) + 2, "Univers d'actifs", span=3)
    asset_header_row = 5 + len(rows) + 3
    ws.cell(row=asset_header_row, column=1, value="Actif")
    ws.cell(row=asset_header_row, column=2, value="Rendement μ (%)")
    ws.cell(row=asset_header_row, column=3, value="Volatilité σ (%)")
    _style_header_row(ws, asset_header_row, 3)

    for i, a in enumerate(req.config.assets, start=asset_header_row + 1):
        ws.cell(row=i, column=1, value=a.name).font = FONT_BODY
        ws.cell(row=i, column=1).alignment = ALIGN_LEFT
        ws.cell(row=i, column=1).border = BORDER_CELL
        ws.cell(row=i, column=2, value=a.expected_return_pct).font = FONT_MONO
        ws.cell(row=i, column=2).alignment = ALIGN_RIGHT
        ws.cell(row=i, column=2).border = BORDER_CELL
        ws.cell(row=i, column=2).number_format = "0.00"
        ws.cell(row=i, column=3, value=a.volatility_pct).font = FONT_MONO
        ws.cell(row=i, column=3).alignment = ALIGN_RIGHT
        ws.cell(row=i, column=3).border = BORDER_CELL
        ws.cell(row=i, column=3).number_format = "0.00"

    _apply_zebra(ws, asset_header_row + 1, asset_header_row + len(req.config.assets), 3)
    ws.column_dimensions["C"].width = 24

    ws.freeze_panes = "A6"


def _build_pareto_sheet(
    wb: Workbook, algo_name: str, algo_data: dict, asset_names: List[str]
) -> None:
    """One row per Pareto-optimal solution."""
    weights_matrix = algo_data.get("weights", [])
    returns = algo_data.get("returns_pct", [])
    risks = algo_data.get("risks_pct", [])
    if not weights_matrix or not returns or not risks:
        return

    ws = wb.create_sheet(f"Pareto_{algo_name}"[:31])

    title = f"Front de Pareto — {algo_name}"
    ws["A1"] = title
    ws["A1"].font = FONT_TITLE
    n_cols = len(asset_names) + 2
    ws.merge_cells(start_row=1, end_row=1, start_column=1, end_column=n_cols)
    ws["A2"] = "Chaque ligne représente un portefeuille non-dominé"
    ws["A2"].font = FONT_CAPTION
    ws.merge_cells(start_row=2, end_row=2, start_column=1, end_column=n_cols)
    ws.row_dimensions[1].height = 28

    header_row = 4
    headers = asset_names + ["Rendement (%)", "Risque (%)"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    _style_header_row(ws, header_row, n_cols)

    n_rows = len(weights_matrix)
    for i, (w_row, ret, risk) in enumerate(
        zip(weights_matrix, returns, risks), start=1
    ):
        r = header_row + i
        for c, w in enumerate(w_row, start=1):
            cell = ws.cell(row=r, column=c, value=w * 100)
            cell.font = FONT_MONO
            cell.alignment = ALIGN_RIGHT
            cell.border = BORDER_CELL
            cell.number_format = '0.00"%"'
        ws.cell(row=r, column=len(asset_names) + 1, value=ret).number_format = '0.00"%"'
        ws.cell(row=r, column=len(asset_names) + 2, value=risk).number_format = (
            '0.00"%"'
        )
        ws.cell(row=r, column=len(asset_names) + 1).font = FONT_MONO
        ws.cell(row=r, column=len(asset_names) + 2).font = FONT_MONO
        ws.cell(row=r, column=len(asset_names) + 1).alignment = ALIGN_RIGHT
        ws.cell(row=r, column=len(asset_names) + 2).alignment = ALIGN_RIGHT
        ws.cell(row=r, column=len(asset_names) + 1).border = BORDER_CELL
        ws.cell(row=r, column=len(asset_names) + 2).border = BORDER_CELL

    _apply_zebra(ws, header_row + 1, header_row + n_rows, n_cols)

    widths = [max(14, len(n) + 2) for n in asset_names] + [16, 14]
    _set_col_widths(ws, widths)
    ws.freeze_panes = f"A{header_row + 1}"


def _build_allocation_sheet(
    wb: Workbook, algo_name: str, algo_data: dict, asset_names: List[str]
) -> None:
    """TOPSIS allocation with data-bar visualization."""
    topsis = algo_data.get("topsis", {})
    if not topsis or not topsis.get("weights"):
        return

    ws = wb.create_sheet(f"Allocation_{algo_name}"[:31])

    ws["A1"] = f"Allocation TOPSIS — {algo_name}"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:C1")
    ws["A2"] = "Compromis optimal sélectionné par arbitrage multi-critères"
    ws["A2"].font = FONT_CAPTION
    ws.merge_cells("A2:C2")
    ws.row_dimensions[1].height = 28

    metrics_row = 4
    _style_section_title(ws, metrics_row, "Métriques de la solution", span=3)

    m_header = metrics_row + 1
    ws.cell(row=m_header, column=1, value="Indicateur")
    ws.cell(row=m_header, column=2, value="Valeur")
    ws.cell(row=m_header, column=3, value="Unité")
    _style_header_row(ws, m_header, 3)

    metrics = [
        ("Rendement attendu (R_p)", round(topsis["return_pct"], 4), "%"),
        ("Volatilité (σ_p)", round(topsis["risk_pct"], 4), "%"),
        ("Ratio de Sharpe", round(topsis["sharpe"], 4), "—"),
    ]
    for i, (lbl, val, unit) in enumerate(metrics, start=m_header + 1):
        ws.cell(row=i, column=1, value=lbl).font = FONT_BODY
        ws.cell(row=i, column=1).alignment = ALIGN_LEFT
        ws.cell(row=i, column=1).border = BORDER_CELL
        ws.cell(row=i, column=2, value=val).font = FONT_MONO
        ws.cell(row=i, column=2).alignment = ALIGN_RIGHT
        ws.cell(row=i, column=2).border = BORDER_CELL
        ws.cell(row=i, column=2).number_format = "0.0000"
        ws.cell(row=i, column=3, value=unit).font = FONT_BODY
        ws.cell(row=i, column=3).alignment = ALIGN_CENTER
        ws.cell(row=i, column=3).border = BORDER_CELL
    _apply_zebra(ws, m_header + 1, m_header + len(metrics), 3)

    alloc_section = m_header + len(metrics) + 2
    _style_section_title(ws, alloc_section, "Répartition par actif", span=3)
    a_header = alloc_section + 1
    ws.cell(row=a_header, column=1, value="Actif")
    ws.cell(row=a_header, column=2, value="Allocation")
    ws.cell(row=a_header, column=3, value="Allocation (%)")
    _style_header_row(ws, a_header, 3)

    sorted_pairs = sorted(
        zip(asset_names, topsis["weights"]),
        key=lambda x: -x[1],
    )

    data_start = a_header + 1
    for i, (name, w) in enumerate(sorted_pairs, start=data_start):
        pct = w * 100
        ws.cell(row=i, column=1, value=name).font = FONT_BODY
        ws.cell(row=i, column=1).alignment = ALIGN_LEFT
        ws.cell(row=i, column=1).border = BORDER_CELL
        ws.cell(row=i, column=2, value=pct)
        ws.cell(row=i, column=2).number_format = "0.00"
        ws.cell(row=i, column=2).font = Font(color="FFFFFF", size=1)
        ws.cell(row=i, column=2).border = BORDER_CELL
        ws.cell(row=i, column=3, value=pct).font = FONT_MONO
        ws.cell(row=i, column=3).alignment = ALIGN_RIGHT
        ws.cell(row=i, column=3).border = BORDER_CELL
        ws.cell(row=i, column=3).number_format = '0.00"%"'

    data_end = data_start + len(sorted_pairs) - 1
    _apply_zebra(ws, data_start, data_end, 3)

    if data_end >= data_start:
        bar_rule = DataBarRule(
            start_type="num",
            start_value=0,
            end_type="num",
            end_value=100,
            color=COLOR_PRIMARY.lstrip("#"),
            showValue=False,
        )
        ws.conditional_formatting.add(
            f"B{data_start}:B{data_end}",
            bar_rule,
        )

    _set_col_widths(ws, [40, 30, 18])
    ws.column_dimensions["B"].width = 30
    ws.freeze_panes = "A5"


def _build_indicators_sheet(wb: Workbook, indicators: Dict[str, Dict]) -> None:
    if not indicators:
        return

    ws = wb.create_sheet("Indicateurs")
    ws["A1"] = "Indicateurs de qualité des fronts"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:C1")
    ws["A2"] = "HV (plus grand = mieux) · IGD (plus petit = mieux)"
    ws["A2"].font = FONT_CAPTION
    ws.merge_cells("A2:C2")
    ws.row_dimensions[1].height = 28

    header_row = 4
    ws.cell(row=header_row, column=1, value="Algorithme")
    ws.cell(row=header_row, column=2, value="Hypervolume")
    ws.cell(row=header_row, column=3, value="IGD")
    _style_header_row(ws, header_row, 3)

    for i, (name, ind) in enumerate(indicators.items(), start=header_row + 1):
        ws.cell(row=i, column=1, value=name).font = FONT_BODY
        ws.cell(row=i, column=1).alignment = ALIGN_LEFT
        ws.cell(row=i, column=1).border = BORDER_CELL
        ws.cell(row=i, column=2, value=ind.get("hv")).font = FONT_MONO
        ws.cell(row=i, column=2).alignment = ALIGN_RIGHT
        ws.cell(row=i, column=2).border = BORDER_CELL
        ws.cell(row=i, column=2).number_format = "0.0000"
        ws.cell(row=i, column=3, value=ind.get("igd")).font = FONT_MONO
        ws.cell(row=i, column=3).alignment = ALIGN_RIGHT
        ws.cell(row=i, column=3).border = BORDER_CELL
        ws.cell(row=i, column=3).number_format = "0.000000"

    _apply_zebra(ws, header_row + 1, header_row + len(indicators), 3)
    _set_col_widths(ws, [22, 18, 18])


@router.post("/export-report")
async def export_report(req: ReportExportRequest):
    if not req.results:
        raise HTTPException(status_code=400, detail="No results to export.")

    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)

    asset_names = [a.name for a in req.config.assets]

    _build_config_sheet(wb, req)

    for algo_name, algo_data in req.results.items():
        if not isinstance(algo_data, dict):
            continue
        _build_pareto_sheet(wb, algo_name, algo_data, asset_names)
        _build_allocation_sheet(wb, algo_name, algo_data, asset_names)

    if req.indicators:
        _build_indicators_sheet(wb, req.indicators)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"Rapport_SID_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
